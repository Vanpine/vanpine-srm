"""
excel 解析器测试

覆盖 P3 验收点：
  (a) 每个 sheet 成一个 heading_path 章节
  (b) 表头保留，每组数据块都重复表头
  (c) 超长表按 TABLE_MAX_ROWS 切片
  (d) 多个 sheet 不混块
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.excel_parser import extract_excel_blocks, TABLE_MAX_ROWS
from app.services.chunker import build_chunks


def _make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "报价单"
    ws1.append(["型号", "数量", "单价"])
    ws1.append(["Q235B", "100吨", "4200"])
    ws1.append(["Q345B", "50吨", "4500"])

    ws2 = wb.create_sheet("交付计划")
    ws2.append(["批次", "交期"])
    ws2.append(["第一批", "2026-03-01"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_long_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["序号", "物料"])
    for i in range(1, TABLE_MAX_ROWS * 2 + 5):
        ws.append([i, f"物料{i}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sheets_as_sections():
    """(a) 每个 sheet 成独立章节。"""
    blocks = extract_excel_blocks(_make_xlsx())
    paths = {tuple(b.heading_path) for b in blocks}
    assert ("报价单",) in paths, f"报价单章节缺失: {paths}"
    assert ("交付计划",) in paths, "交付计划章节缺失"


def test_header_preserved():
    """(b) 表头保留。"""
    blocks = extract_excel_blocks(_make_xlsx())
    md = blocks[0].text
    assert "| 型号 | 数量 | 单价 |" in md, f"表头缺失: {md}"
    assert "Q235B" in md and "4200" in md


def test_long_sheet_split_with_repeated_header():
    """(c) 超长表切片且每片重复表头。"""
    blocks = extract_excel_blocks(_make_long_xlsx())
    assert len(blocks) >= 2, "超长表未切片"
    for b in blocks:
        assert b.text.startswith("| 序号 | 物料 |"), "后续片缺表头"


def test_sheets_not_mixed():
    """(d) 不同 sheet 的内容不混在同一块。"""
    blocks = extract_excel_blocks(_make_xlsx())
    for b in blocks:
        assert not ("Q235B" in b.text and "第一批" in b.text), "跨 sheet 混块"


if __name__ == "__main__":
    import traceback
    passed = failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"[PASS] {name}")
                passed += 1
            except Exception as e:
                print(f"[FAIL] {name}: {e}")
                traceback.print_exc()
                failed += 1
    print(f"\n{passed} passed, {failed} failed")
    sys.exit(1 if failed else 0)
